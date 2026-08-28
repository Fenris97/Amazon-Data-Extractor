"""Read and validate Amazon ASINs from Excel workbooks."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


ASIN_PATTERN = re.compile(r"^[A-Z0-9]{10}$")
AMAZON_URL_PATTERNS = (
    re.compile(r"/(?:dp|gp/product|product)/([A-Z0-9]{10})(?:[/?#]|$)", re.I),
    re.compile(r"[?&]asin=([A-Z0-9]{10})(?:[&#]|$)", re.I),
)
DEFAULT_COLUMN_ALIASES = (
    "asin",
    "child asin",
    "amazon asin",
    "product asin",
)


class InputWorkbookError(ValueError):
    """Raised when the workbook cannot provide a usable ASIN column."""


@dataclass(frozen=True)
class AsinAuditRow:
    source_row: int
    original_value: str
    normalized_asin: str
    status: str
    message: str


@dataclass(frozen=True)
class AsinReadResult:
    asins: list[str]
    audit_rows: list[AsinAuditRow]
    sheet_name: str
    column_name: str

    @property
    def invalid_count(self) -> int:
        return sum(row.status == "INVALID" for row in self.audit_rows)

    @property
    def duplicate_count(self) -> int:
        return sum(row.status == "DUPLICATE" for row in self.audit_rows)


def _normalize_header(value: object) -> str:
    return " ".join(str(value or "").strip().lower().replace("_", " ").split())


def extract_asin(value: object) -> str | None:
    """Return a normalized ASIN from a cell value or Amazon URL."""

    if value is None:
        return None

    text = str(value).strip()
    if not text:
        return None

    normalized = text.upper()
    if ASIN_PATTERN.fullmatch(normalized):
        return normalized

    for pattern in AMAZON_URL_PATTERNS:
        match = pattern.search(text)
        if match:
            return match.group(1).upper()

    return None


def _find_column(headers: Iterable[object], requested_column: str | None) -> tuple[int, str]:
    header_list = list(headers)
    normalized_headers = [_normalize_header(value) for value in header_list]

    candidates = (requested_column,) if requested_column else DEFAULT_COLUMN_ALIASES
    for candidate in candidates:
        normalized_candidate = _normalize_header(candidate)
        if normalized_candidate in normalized_headers:
            index = normalized_headers.index(normalized_candidate)
            return index, str(header_list[index]).strip()

    available = [str(value).strip() for value in header_list if value not in (None, "")]
    requested = requested_column or "ASIN / Child ASIN"
    raise InputWorkbookError(
        f"Could not find column '{requested}'. Available columns: "
        + (", ".join(available) if available else "none")
    )


def read_asins(
    workbook_path: str | Path,
    *,
    sheet_name: str | None = None,
    column_name: str | None = None,
    header_row: int = 1,
) -> AsinReadResult:
    """Read unique valid ASINs and return an audit trail for populated rows."""

    path = Path(workbook_path)
    if not path.is_file():
        raise InputWorkbookError(f"Excel file not found: {path}")
    if path.suffix.lower() != ".xlsx":
        raise InputWorkbookError("The input must be an .xlsx Excel workbook.")
    if header_row < 1:
        raise InputWorkbookError("Header row must be 1 or greater.")

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise InputWorkbookError(f"Could not open the Excel workbook: {exc}") from exc

    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise InputWorkbookError(
                    f"Worksheet '{sheet_name}' was not found. Available worksheets: "
                    + ", ".join(workbook.sheetnames)
                )
            worksheet = workbook[sheet_name]
        else:
            worksheet = workbook.active

        header_values = next(
            worksheet.iter_rows(
                min_row=header_row,
                max_row=header_row,
                values_only=True,
            ),
            (),
        )
        column_index, actual_column_name = _find_column(header_values, column_name)

        accepted: list[str] = []
        audit_rows: list[AsinAuditRow] = []
        seen: set[str] = set()

        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            if not any(value not in (None, "") for value in row):
                continue

            value = row[column_index] if column_index < len(row) else None
            original = "" if value is None else str(value).strip()
            asin = extract_asin(value)

            if not original:
                audit_rows.append(
                    AsinAuditRow(row_number, original, "", "INVALID", "ASIN cell is blank")
                )
            elif asin is None:
                audit_rows.append(
                    AsinAuditRow(
                        row_number,
                        original,
                        "",
                        "INVALID",
                        "Expected a 10-character ASIN or Amazon product URL",
                    )
                )
            elif asin in seen:
                audit_rows.append(
                    AsinAuditRow(
                        row_number,
                        original,
                        asin,
                        "DUPLICATE",
                        "Already accepted from an earlier row",
                    )
                )
            else:
                seen.add(asin)
                accepted.append(asin)
                audit_rows.append(AsinAuditRow(row_number, original, asin, "ACCEPTED", ""))

        return AsinReadResult(
            asins=accepted,
            audit_rows=audit_rows,
            sheet_name=worksheet.title,
            column_name=actual_column_name,
        )
    finally:
        workbook.close()

