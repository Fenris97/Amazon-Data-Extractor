"""Write the untouched JSON response and a readable Excel workbook."""

from __future__ import annotations

import json
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .asin_reader import AsinAuditRow


# Stable field order from the confirmed Bright Data Amazon product schema.
# Unknown fields are appended, so schema additions are not discarded.
KNOWN_FIELDS = (
    "category_tree", "target_countries", "category_urls", "title", "seller_name",
    "brand", "description", "initial_price", "currency", "availability",
    "reviews_count", "categories", "parent_asin", "asin", "buybox_seller",
    "number_of_sellers", "root_bs_rank", "ISBN10", "answered_questions", "domain",
    "images_count", "url", "video_count", "image_url", "item_weight", "rating",
    "product_dimensions", "seller_id", "image", "date_first_available", "discount",
    "model_number", "manufacturer", "department", "plus_content", "upc", "video",
    "top_review", "final_price_high", "final_price", "variations", "delivery",
    "features", "format", "buybox_prices", "input_asin", "ingredients", "origin_url",
    "bought_past_month", "is_available", "root_bs_category", "bs_category", "bs_rank",
    "badge", "subcategory_rank", "amazon_choice", "images", "product_details",
    "prices_breakdown", "country_of_origin", "from_the_brand", "product_description",
    "seller_url", "customer_says", "sustainability_features", "climate_pledge_friendly",
    "videos", "other_sellers_prices", "downloadable_videos", "editorial_reviews",
    "about_the_author", "zipcode", "coupon", "sponsered", "store_url", "ships_from",
    "city", "customers_say", "max_quantity_available", "variations_values", "language",
    "return_policy", "inactive_buy_box", "buybox_seller_rating", "premium_brand",
    "amazon_prime", "coupon_description", "all_badges", "sponsored",
    "variant_attributes", "safety_information", "subcategory_link",
    "is_frequently_returned_item_badge", "frequently_returned_item_message",
    "is_customers_usually_keep", "title_badge", "review_images", "review_videos",
    "bought_past_month_text", "is_high_price", "title_highlight", "title_clean",
    "customers_say_topics", "timestamp", "input_url", "input_language",
    "input_all_variations", "input_zipcode", "error", "error_code", "warning",
    "warning_code",
)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
EXCEL_CELL_LIMIT = 32_767


def _json_default(value: object) -> str:
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)


def _excel_value(value: Any) -> str | int | float | bool:
    if value is None:
        return ""
    if isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, (dict, list, tuple)):
        text = json.dumps(value, ensure_ascii=False, default=_json_default)
    else:
        text = str(value)

    if len(text) > EXCEL_CELL_LIMIT:
        text = text[:32_680] + "... [complete value is in brightdata_raw.json]"
    if text.startswith(("=", "+", "-", "@")):
        text = "'" + text
    return text


def _ordered_fields(records: Iterable[dict[str, Any]]) -> list[str]:
    discovered: set[str] = set()
    for record in records:
        discovered.update(record)
    known = list(KNOWN_FIELDS)
    unknown = sorted(discovered.difference(KNOWN_FIELDS), key=str.lower)
    return known + unknown


def _style_table(worksheet, *, freeze: str = "A2", max_width: int = 50) -> None:
    worksheet.freeze_panes = freeze
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for column_cells in worksheet.iter_cols():
        letter = get_column_letter(column_cells[0].column)
        width = min(
            max((len(str(cell.value or "")) for cell in column_cells[:100]), default=8) + 2,
            max_width,
        )
        worksheet.column_dimensions[letter].width = max(10, width)


def write_outputs(
    run_directory: str | Path,
    *,
    records: list[dict[str, Any]],
    audit_rows: list[AsinAuditRow],
    metadata: dict[str, Any],
) -> tuple[Path, Path, Path]:
    """Write raw JSON, metadata, and the organized workbook."""

    run_dir = Path(run_directory)
    run_dir.mkdir(parents=True, exist_ok=True)
    raw_json_path = run_dir / "brightdata_raw.json"
    metadata_path = run_dir / "run_metadata.json"
    workbook_path = run_dir / "amazon_extraction.xlsx"

    raw_json_path.write_text(
        json.dumps(records, indent=2, ensure_ascii=False, default=_json_default),
        encoding="utf-8",
    )
    metadata_path.write_text(
        json.dumps(metadata, indent=2, ensure_ascii=False, default=_json_default),
        encoding="utf-8",
    )

    workbook = Workbook()
    run_sheet = workbook.active
    run_sheet.title = "Run Information"
    run_sheet.append(["Field", "Value"])
    for key, value in metadata.items():
        run_sheet.append([key, _excel_value(value)])
    _style_table(run_sheet, max_width=70)

    data_sheet = workbook.create_sheet("Amazon Data")
    fields = _ordered_fields(records)
    data_sheet.append(fields)
    for record in records:
        data_sheet.append([_excel_value(record.get(field)) for field in fields])
    _style_table(data_sheet)
    data_sheet.sheet_view.showGridLines = False

    audit_sheet = workbook.create_sheet("Input Audit")
    audit_headers = ["source_row", "original_value", "normalized_asin", "status", "message"]
    audit_sheet.append(audit_headers)
    for audit_row in audit_rows:
        values = asdict(audit_row)
        audit_sheet.append([_excel_value(values[header]) for header in audit_headers])
    _style_table(audit_sheet)

    error_sheet = workbook.create_sheet("Errors")
    error_headers = (
        "input_url", "input_asin", "asin", "error", "error_code", "warning", "warning_code"
    )
    error_sheet.append(error_headers)
    for record in records:
        if record.get("error") or record.get("warning"):
            error_sheet.append([_excel_value(record.get(field)) for field in error_headers])
    _style_table(error_sheet)

    workbook.save(workbook_path)
    workbook.close()
    return raw_json_path, workbook_path, metadata_path

