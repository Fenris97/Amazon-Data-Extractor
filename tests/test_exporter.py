import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from amazon_data_extractor.asin_reader import AsinAuditRow
from amazon_data_extractor.exporter import write_outputs


class ExporterTests(unittest.TestCase):
    def test_writes_raw_json_and_expected_excel_sheets(self):
        records = [
            {
                "asin": "B0G38V6MRN",
                "title": "Example product",
                "features": ["First bullet", "Second bullet"],
                "new_future_field": {"enabled": True},
            }
        ]
        audit = [AsinAuditRow(2, "B0G38V6MRN", "B0G38V6MRN", "ACCEPTED", "")]

        with tempfile.TemporaryDirectory() as directory:
            raw_path, workbook_path, metadata_path = write_outputs(
                directory,
                records=records,
                audit_rows=audit,
                metadata={"snapshot_id": "s_test", "status": "completed"},
            )

            self.assertEqual(json.loads(raw_path.read_text(encoding="utf-8")), records)
            self.assertTrue(metadata_path.is_file())

            workbook = load_workbook(workbook_path, read_only=True)
            self.assertEqual(
                workbook.sheetnames,
                ["Run Information", "Amazon Data", "Input Audit", "Errors"],
            )
            data_sheet = workbook["Amazon Data"]
            headers = [cell.value for cell in next(data_sheet.iter_rows(max_row=1))]
            self.assertIn("new_future_field", headers)
            feature_column = headers.index("features") + 1
            self.assertEqual(
                data_sheet.cell(row=2, column=feature_column).value,
                '["First bullet", "Second bullet"]',
            )
            workbook.close()


if __name__ == "__main__":
    unittest.main()

